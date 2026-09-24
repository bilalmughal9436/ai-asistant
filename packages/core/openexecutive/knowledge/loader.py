from __future__ import annotations

import contextlib
import hashlib
import logging
import re
from pathlib import Path
from typing import Any

from openexecutive.knowledge.store import ChromaDBStore

logger = logging.getLogger(__name__)

BUILTIN_KNOWLEDGE_PATH = Path(__file__).parent / "builtin"
FAILURES_KNOWLEDGE_PATH = BUILTIN_KNOWLEDGE_PATH / "failures"

DOMAIN_MAP: dict[str, str] = {
    "strategy": "strategy",
    "finance": "finance",
    "hr": "hr",
    "legal": "legal",
    "operations": "operations",
    "marketing": "marketing",
    "board": "board",
    "product": "product",
}

# The catch-all domain for company documents the uploader didn't classify.
# Unlike the eight specialist domains it maps to no single specialist — every
# specialist retrieves it (see ``retriever.retrieve``), so an unclassified
# upload is visible to all rather than to none.
GENERAL_DOMAIN = "general"

# Domains ``POST /documents`` accepts. Anything else is a typo that would
# silently index the document where no specialist can ever retrieve it.
UPLOAD_DOMAINS: frozenset[str] = frozenset(DOMAIN_MAP) | {GENERAL_DOMAIN}

# Chunk-id namespace for a file attached in an integration channel. The
# sender picks the filename and the filename is the id namespace, so an
# unprefixed ``strategy-2026.md`` would upsert over whatever else carries
# that name. Shared between the write path
# (``integrations.attachments``) and the migration below, which has to
# recognise exactly what the write path produced.
ATTACHMENT_SOURCE_PREFIX = "attachment:"

# Domain tag on attachment chunks. Isolation is the collection
# (``ChromaDBStore.ATTACHMENT_COLLECTION``), never this value — but if
# anything ever does query that collection, this keeps the default
# "not retrieved": it is in neither UPLOAD_DOMAINS nor DOMAIN_ALIASES, and
# deliberately NOT ``general``, which ``retriever._with_general`` would fan
# out to every specialist.
ATTACHMENT_DOMAIN = "attachment"

# What the pre-isolation attachment path wrote into ``domain``. Only the
# boot migration reads it, to tell its own legacy rows from a curated
# upload that merely happens to be named with the prefix. UPLOAD_DOMAINS
# never accepts this value, so nothing uploaded through the API carries it.
_LEGACY_ATTACHMENT_DOMAIN = "company_docs"


def chunk_text(text: str, chunk_size: int = 512, overlap: int = 50) -> list[str]:
    words = text.split()
    if not words:
        return []

    chunks: list[str] = []
    start = 0
    while start < len(words):
        end = min(start + chunk_size, len(words))
        chunk = " ".join(words[start:end])
        chunks.append(chunk)
        if end == len(words):
            break
        start = end - overlap
    return chunks


def extract_text_from_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text.strip())
    return "\n\n".join(pages)


def extract_text_from_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def extract_text_from_xlsx(path: Path, max_chars: int = 200_000) -> str:
    """Flatten an .xlsx/.xlsm workbook to text — one ``## <sheet>`` heading per
    NON-EMPTY worksheet, cells tab-joined and rows newline-joined. Only reads
    stored cell values (``data_only=True`` returns cached formula results, not
    formulae); legacy binary ``.xls`` is not supported by openpyxl.

    ``read_only`` streams rows and ``max_chars`` bounds the accumulated text, so
    a decompression-bombed workbook (a small archive that inflates to millions
    of cells) can't exhaust memory."""
    from openpyxl import load_workbook

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        total = 0
        for ws in wb.worksheets:
            heading_written = False
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if not cells:
                    continue
                if not heading_written:
                    # Defer the heading until the sheet is known to have data,
                    # so a fully-blank sheet contributes nothing.
                    heading = f"## {ws.title}"
                    parts.append(heading)
                    total += len(heading) + 1
                    heading_written = True
                line = "\t".join(cells)
                parts.append(line)
                total += len(line) + 1
                if total >= max_chars:
                    return "\n".join(parts)
        return "\n".join(parts)
    finally:
        wb.close()


def extract_text_from_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(path)
    elif suffix in (".docx", ".doc"):
        return extract_text_from_docx(path)
    elif suffix in (".xlsx", ".xlsm"):
        return extract_text_from_xlsx(path)
    elif suffix in (".md", ".txt", ".rst", ".csv"):
        return path.read_text(encoding="utf-8")
    return ""


def _make_chunk_id(source: str, chunk_index: int) -> str:
    base = f"{source}::chunk::{chunk_index}"
    return hashlib.md5(base.encode()).hexdigest()


def infer_domain_from_path(path: Path, root: Path | None = None) -> str:
    """Infer a knowledge domain from a path's components.

    Pass ``root`` for content under a known tree: only the parts BELOW it are
    scanned. Without it this walks the ABSOLUTE path, so an install rooted at
    e.g. `/srv/product/` tags every chunk `product` no matter which domain
    directory the file is actually in. That was merely a mis-tag until the
    review gate began keying on `(domain, filename)` — a chunk domain that
    disagrees with its review row's domain means the key never matches and
    withheld content is retrieved. Callers under `knowledge/builtin/` must
    pass the root so the metadata agrees with what `review_items` stores.
    """
    parts = path.parts
    if root is not None:
        # A path outside the root keeps its absolute parts — the caller asked
        # for a hint, not a constraint.
        with contextlib.suppress(ValueError):
            parts = path.relative_to(root).parts
    for part in parts:
        domain = DOMAIN_MAP.get(part.lower())
        if domain:
            return domain
    return "general"


async def ingest_file(
    path: Path,
    store: ChromaDBStore,
    domain: str | None = None,
    collection: str = ChromaDBStore.COMPANY_COLLECTION,
    *,
    source_name: str | None = None,
    extra_metadata: dict[str, Any] | None = None,
) -> int:
    """Index a file on disk into a knowledge collection.

    ``source_name`` is the *logical* identity of the document — it becomes the
    ``filename``/``source`` metadata and the chunk-id namespace, exactly as in
    ``ingest_text_sync``. Callers that stage an upload through a temp file MUST
    pass the real filename: ``path`` is then a random ``tmpXXXXXXXX`` name, and
    deriving identity from it gives every re-upload fresh chunk ids (so the
    id-keyed upsert never collides and duplicates without bound) and stores a
    ``filename`` that ``DELETE /documents/{filename}`` can never match.

    Defaults to ``path.name`` rather than ``str(path)`` so that identity is the
    bare filename for every caller. That keeps an uploaded ``plan.md`` and a
    fixture-loaded ``plan.md`` on the same ids instead of duplicating each
    other, and matches how ``DELETE`` and ``list_company_docs`` already treat
    filename as the document's identity.

    ``extra_metadata`` is merged into every chunk, mirroring
    ``ingest_text_sync``. It is how a collection gets a ``type`` tag it can
    later be deleted by: Chroma's ``where`` matches exact values only, so a
    tag is the difference between a one-call delete and a full metadata scan.
    """
    text = extract_text_from_file(path)
    if not text.strip():
        return 0

    name = source_name or path.name
    inferred_domain = domain or infer_domain_from_path(path)
    chunks = chunk_text(text, chunk_size=512, overlap=50)

    texts = chunks
    metadatas: list[dict[str, Any]] = [
        {
            "domain": inferred_domain,
            "filename": name,
            "source": name,
            "chunk_index": i,
            **(extra_metadata or {}),
        }
        for i in range(len(chunks))
    ]
    ids = [_make_chunk_id(name, i) for i in range(len(chunks))]

    store.add_documents(texts=texts, metadatas=metadatas, ids=ids, collection=collection)
    return len(chunks)


def ingest_text_sync(
    text: str,
    store: ChromaDBStore,
    *,
    source_name: str,
    domain: str = "general",
    collection: str = ChromaDBStore.COMPANY_COLLECTION,
    extra_metadata: dict[str, Any] | None = None,
) -> int:
    """Synchronous ingest of a raw markdown/text string.

    Mirrors ``ingest_file`` but takes a string — used to persist research
    artifacts and Notion wiki pages into a named collection. ``source_name``
    is the logical identifier for ``filename``/``source`` metadata and the
    chunk-id namespace. ``extra_metadata`` is merged into every chunk.
    Returns the number of chunks written.

    Callers on the API event loop should wrap this in ``asyncio.to_thread``.
    """
    if not text.strip():
        return 0

    chunks = chunk_text(text, chunk_size=512, overlap=50)
    extra = extra_metadata or {}
    metadatas: list[dict[str, Any]] = [
        {
            "domain": domain,
            "filename": source_name,
            "source": source_name,
            "chunk_index": i,
            **extra,
        }
        for i in range(len(chunks))
    ]
    ids = [_make_chunk_id(source_name, i) for i in range(len(chunks))]

    store.add_documents(texts=chunks, metadatas=metadatas, ids=ids, collection=collection)
    return len(chunks)


async def ingest_text(
    text: str,
    store: ChromaDBStore,
    *,
    source_name: str,
    domain: str = "general",
    collection: str = ChromaDBStore.COMPANY_COLLECTION,
    extra_metadata: dict[str, Any] | None = None,
) -> int:
    """Ingest a raw markdown/text string as knowledge (no file on disk).

    Mirrors ``ingest_file`` but takes a string — used to persist the
    executive_research artifact into its own collection. ``source_name``
    is the logical identifier used for both the ``filename``/``source``
    metadata and the chunk-id namespace. ``extra_metadata`` is merged into
    every chunk's metadata (e.g. ``{"type": "recent_research", "created_at": …}``).
    Returns the number of chunks written.
    """
    return ingest_text_sync(
        text,
        store,
        source_name=source_name,
        domain=domain,
        collection=collection,
        extra_metadata=extra_metadata,
    )


async def ingest_builtin_file(
    path: Path,
    store: ChromaDBStore,
    collection: str = ChromaDBStore.BUILTIN_COLLECTION,
    chunk_type: str = "builtin",
    chunk_size: int = 512,
    overlap: int = 50,
) -> int:
    """Index a single built-in markdown file. Caller must delete old chunks first.

    Defaults match the positive-playbook ingest path. Pass
    ``collection=ChromaDBStore.FAILURES_COLLECTION`` (with ``chunk_type='failure_case'``
    and smaller chunks) to ingest a single failure case study — keeps the
    failure CRUD endpoints in lockstep with ``seed_failures``.
    """
    text = path.read_text(encoding="utf-8")
    if not text.strip():
        return 0
    domain = infer_domain_from_path(path, root=BUILTIN_KNOWLEDGE_PATH)
    chunks = chunk_text(text, chunk_size=chunk_size, overlap=overlap)
    metadatas: list[dict[str, Any]] = [
        {
            "domain": domain,
            "filename": path.name,
            "source": str(path),
            "chunk_index": i,
            "type": chunk_type,
        }
        for i in range(len(chunks))
    ]
    ids = [_make_chunk_id(str(path), i) for i in range(len(chunks))]
    store.add_documents(
        texts=chunks,
        metadatas=metadatas,
        ids=ids,
        collection=collection,
    )
    return len(chunks)


def _company_doc_names(docs_dir: Path) -> set[str]:
    """Filenames in ``docs_dir`` that count as company documents.

    One definition, shared by the listing endpoint and the boot reconcile, so
    the two can never disagree about what is on disk. Excludes subdirectories
    (Notion sync writes into ``docs/notion/``, which belongs to its own
    collection) and dotfiles.
    """
    if not docs_dir.is_dir():
        return set()
    return {p.name for p in docs_dir.iterdir() if p.is_file() and not p.name.startswith(".")}


def list_company_docs(docs_dir: Path) -> list[dict[str, Any]]:
    if not docs_dir.exists():
        return []
    return [
        {
            "filename": name,
            "size_bytes": (docs_dir / name).stat().st_size,
            "modified_at": (docs_dir / name).stat().st_mtime,
        }
        for name in sorted(_company_doc_names(docs_dir))
    ]


# Names produced by ``tempfile.NamedTemporaryFile(suffix=ext)`` — "tmp" plus a
# random stem, e.g. "tmp1du9epr4.md". Uploads used to be indexed under these
# (the staging path was passed to ``ingest_file`` instead of the real name), so
# deployed stores hold chunks no filename can ever match.
_TEMP_CHUNK_NAME = re.compile(r"^tmp[A-Za-z0-9_]{6,12}\.[A-Za-z0-9]{1,5}$")


async def reconcile_company_docs(
    store: ChromaDBStore,
    docs_dir: Path,
    collection: str = ChromaDBStore.COMPANY_COLLECTION,
) -> tuple[int, int]:
    """Sweep temp-named orphans, then index documents that have no chunks.

    Returns ``(orphans_deleted, files_indexed)``. Converges on a stable store,
    so it is safe to run on every boot.

    Only files with *no* rows in the collection are indexed. A document that is
    already indexed under its real filename is left untouched — its ``domain``
    came from whatever the uploader chose, and the on-disk copy carries no
    record of that, so re-ingesting would silently retag a ``finance`` document
    as ``general``. Recovered documents land under ``general`` explicitly (NOT
    via ``infer_domain_from_path``, which scans every component of the absolute
    path and would tag them ``finance`` on an install rooted under, say,
    ``/srv/finance/``). That is the honest answer — their domain died with the
    orphaned rows — and ``general`` is retrievable by every specialist rather
    than by none.

    Documents whose only copy was the orphaned index — chat/email attachments,
    which are never written to ``docs_dir`` — cannot be recovered and are
    dropped by the sweep.

    A document that extracts to no text writes no chunks, so it is retried on
    every boot. That is wasted work rather than churn (the store still
    converges), and it is bounded by however many unreadable files are sitting
    in ``docs_dir``.

    Does nothing at all when ``docs_dir`` does not exist. An absent directory is
    ambiguous — "no documents" and "the volume is not mounted yet" look
    identical — and in the second case every recoverable document would be
    swept precisely because the file that would have spared it is invisible.
    """
    if not docs_dir.is_dir():
        logger.warning(
            "reconcile_company_docs: %s does not exist — skipping "
            "(cannot distinguish an empty docs dir from an unmounted one)",
            docs_dir,
        )
        return 0, 0

    live = _company_doc_names(docs_dir)

    orphans: list[str] = []
    indexed_names: set[str] = set()
    for cid, md in store.iter_chunk_metadata(collection):
        name = md.get("filename")
        if not isinstance(name, str):
            continue
        if _TEMP_CHUNK_NAME.match(name) and name not in live:
            orphans.append(cid)
        else:
            indexed_names.add(name)

    deleted = store.delete_by_ids(collection, orphans)

    indexed = 0
    for doc_name in sorted(live - indexed_names):
        try:
            if await ingest_file(
                docs_dir / doc_name,
                store,
                domain=GENERAL_DOMAIN,
                collection=collection,
            ):
                indexed += 1
        except Exception:
            # One unreadable document must not abort the reconcile —
            # the rest of the index is still worth repairing.
            logger.exception("reconcile_company_docs: index failed: %s", doc_name)

    return deleted, indexed


def _is_attachment_row(metadata: dict[str, Any]) -> bool:
    """True when a COMPANY row was written by the attachment path.

    The prefix alone is not enough. A colon is legal in a filename and
    ``POST /documents`` passes one through, so a curated upload can be called
    ``attachment:q4-plan.md`` — matching on the name alone would migrate a
    real company document out of the collection it belongs in, and (because
    the file is still on disk) the boot reconcile would re-index it as
    ``general`` on the way back, widening it to every specialist. So the name
    has to be corroborated by ``domain="company_docs"``, the value the old
    attachment path always wrote and the one value ``UPLOAD_DOMAINS`` will
    never accept from an upload.

    ``type="attachment"`` is checked first as defence in depth: this change
    introduces the tag, so no shipped build ever wrote one into COMPANY, but
    a row carrying it is unambiguously ours.

    ``startswith``, not ``in`` — ``attachments-policy.md`` is a document about
    attachments, not an attachment.
    """
    if metadata.get("type") == "attachment":
        return True
    if metadata.get("domain") != _LEGACY_ATTACHMENT_DOMAIN:
        return False
    return any(
        str(metadata.get(key) or "").startswith(ATTACHMENT_SOURCE_PREFIX)
        for key in ("filename", "source")
    )


def migrate_attachments_out_of_company_docs(store: ChromaDBStore) -> int:
    """Move attachment chunks from COMPANY into the isolated collection.

    Attachment rows used to be written into ``company_docs`` and kept out of
    retrieval by a domain outside the specialist set, which never excluded
    them (see knowledge.general_catch_all in architecture-facts.yaml).
    Isolation is now the collection, and the rows already in COMPANY have to
    follow.

    Returns the number of chunks actually removed from COMPANY — not the
    number attempted. A row that was copied but not deleted is still
    retrievable, so reporting it as moved would announce the leak as closed
    while it is open.

    Idempotent: ids are reused verbatim, an id already present in the
    destination is left alone, and once a run completes the scan that
    precedes it finds nothing.

    Synchronous on purpose. Every step here blocks — the metadata scan, the
    text fetch, the embedding pass inside ``add_documents`` — so the caller
    has to put it on a thread rather than merely scheduling it on the event
    loop, which would defer the stall without avoiding it.

    A note for anyone reading a zero on an install that certainly received
    attachments: before this change the ingest path built its store with the
    bare ``ChromaDBStore()`` default rather than ``settings.vector_store_path``,
    so wherever those differ the rows went to a different database than the
    one the API reads. In a container that directory is outside the data
    volume and goes with the container; elsewhere it persists as a stray
    ``<cwd>/chroma_db`` that no wipe path reaches and an operator should
    remove by hand. Either way it is unreachable from here, and zero is the
    honest answer rather than a failure.
    """
    candidates = [
        chunk_id
        for chunk_id, metadata in store.iter_chunk_metadata(
            ChromaDBStore.COMPANY_COLLECTION
        )
        if _is_attachment_row(metadata)
    ]
    if not candidates:
        return 0

    rows = store.get_documents_by_ids(ChromaDBStore.COMPANY_COLLECTION, candidates)
    if not rows:
        # The scan found rows and the fetch returned none. Those two cannot
        # both be right, and `get_documents_by_ids` reports a read failure the
        # same way it reports an empty result — so say so rather than letting
        # a broken read look like a clean store.
        logger.error(
            "attachment migration: %d row(s) matched but none could be read "
            "from %s; leaving them in place",
            len(candidates),
            ChromaDBStore.COMPANY_COLLECTION,
        )
        return 0

    # Ids collide by design: they are md5 of the same source name on both
    # sides. A row already in the destination is a NEWER copy written by the
    # post-fix path, so upserting the COMPANY copy over it would restore
    # superseded text. Skip those and just drop the stale original.
    already_there = {
        chunk_id
        for chunk_id, _, _ in store.get_documents_by_ids(
            ChromaDBStore.ATTACHMENT_COLLECTION, [cid for cid, _, _ in rows]
        )
    }
    fresh = [row for row in rows if row[0] not in already_there]

    if fresh:
        ids = [chunk_id for chunk_id, _, _ in fresh]
        texts = [text for _, text, _ in fresh]
        metadatas: list[dict[str, Any]] = []
        for _, _, metadata in fresh:
            # Normalise onto one shape. Nothing reads these values today —
            # the collection is never queried — which is why it is cheap now
            # and expensive once something does.
            moved = dict(metadata)
            moved["type"] = "attachment"
            moved["domain"] = ATTACHMENT_DOMAIN
            metadatas.append(moved)

        # Write first, delete second, and never the other way round. Chroma
        # spans no transaction across two collections: a crash between them
        # leaves a duplicate the next run collapses on unchanged ids, while
        # deleting first would destroy the only copy — an attachment is never
        # written to ``company/docs/``, so nothing can re-ingest it.
        store.add_documents(
            texts=texts,
            metadatas=metadatas,
            ids=ids,
            collection=ChromaDBStore.ATTACHMENT_COLLECTION,
        )

    stale_ids = [chunk_id for chunk_id, _, _ in rows]
    deleted = store.delete_by_ids(ChromaDBStore.COMPANY_COLLECTION, stale_ids)
    if deleted != len(stale_ids):
        logger.error(
            "attachment migration: copied %d row(s) but removed %d from %s — "
            "the remainder are still retrievable",
            len(stale_ids),
            deleted,
            ChromaDBStore.COMPANY_COLLECTION,
        )
    return deleted


async def seed_builtin_knowledge(
    store: ChromaDBStore | None = None,
    force: bool = False,
) -> int:
    if store is None:
        from openexecutive.config import get_settings

        settings = get_settings()
        store = ChromaDBStore(persist_directory=settings.vector_store_path)

    if not force and store.get_collection_count(ChromaDBStore.BUILTIN_COLLECTION) > 0:
        return 0

    total = 0
    for md_file in BUILTIN_KNOWLEDGE_PATH.rglob("*.md"):
        # Skills live under builtin/skills/ but are indexed into a separate
        # collection by openexecutive.knowledge.skills_index.seed_builtin_skills.
        if any(p in md_file.relative_to(BUILTIN_KNOWLEDGE_PATH).parts for p in ("skills", "failures")):
            continue
        domain = infer_domain_from_path(md_file, root=BUILTIN_KNOWLEDGE_PATH)
        text = md_file.read_text(encoding="utf-8")
        chunks = chunk_text(text)

        metadatas: list[dict[str, Any]] = [
            {
                "domain": domain,
                "filename": md_file.name,
                "source": str(md_file),
                "chunk_index": i,
                "type": "builtin",
            }
            for i in range(len(chunks))
        ]
        ids = [_make_chunk_id(str(md_file), i) for i in range(len(chunks))]
        store.add_documents(
            texts=chunks,
            metadatas=metadatas,
            ids=ids,
            collection=ChromaDBStore.BUILTIN_COLLECTION,
        )
        total += len(chunks)

    return total


async def seed_failures(
    store: ChromaDBStore | None = None,
    force: bool = False,
) -> int:
    """Index all failure case studies from builtin/failures/<domain>/*.md.

    Idempotent: skipped if the failures collection is already non-empty,
    unless force=True. Uses a smaller chunk size (400 words) to preserve
    the narrative arc of each section (situation/root-cause/lessons).
    """
    if store is None:
        from openexecutive.config import get_settings

        settings = get_settings()
        store = ChromaDBStore(persist_directory=settings.vector_store_path)

    if not force and store.get_collection_count(ChromaDBStore.FAILURES_COLLECTION) > 0:
        return 0

    if not FAILURES_KNOWLEDGE_PATH.is_dir():
        logger.warning("failures knowledge path not found, skipping: %s", FAILURES_KNOWLEDGE_PATH)
        return 0

    total = 0
    for md_file in FAILURES_KNOWLEDGE_PATH.rglob("*.md"):
        domain = infer_domain_from_path(md_file, root=BUILTIN_KNOWLEDGE_PATH)
        text = md_file.read_text(encoding="utf-8")
        if not text.strip():
            continue
        chunks = chunk_text(text, chunk_size=400, overlap=40)
        metadatas: list[dict[str, Any]] = [
            {
                "domain": domain,
                "filename": md_file.name,
                "source": str(md_file),
                "chunk_index": i,
                "type": "failure_case",
            }
            for i in range(len(chunks))
        ]
        ids = [_make_chunk_id(str(md_file), i) for i in range(len(chunks))]
        store.add_documents(
            texts=chunks,
            metadatas=metadatas,
            ids=ids,
            collection=ChromaDBStore.FAILURES_COLLECTION,
        )
        total += len(chunks)

    return total
